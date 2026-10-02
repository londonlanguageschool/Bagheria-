"""Build lesson-plans.js from the owner's PowerPoint decks.

Usage:  python3 tools/build-lesson-plans.py /path/to/decks > lesson-plans.js

Folder layout: /path/to/decks/<Folder>/<Prefix>_<CODE>_<Title>.pptx
  CODE = 1B, 10A-1 (two lessons for one book lesson), 1-2 (Revise and Check
  units 1-2, becomes R1-2) or L01 (stand-alone course lesson, becomes L1).
Slide 1 texts: code, title, focus ("Grammar • vocabulary"), question,
  "Student's Book pp.8–9" (optional). Other slides: text 2 is the section
  label (Reading, Listening, Speaking, Vocabulary, Grammar, …game).
Add a new course: put its folder in COURSES with the portal Book code.
"""
import json, os, re, sys
from pptx import Presentation

COURSES = {  # folder -> (Book code used in the Classes sheet, course name)
    "A2_Blue_Part1_Units1-6": ("EF-A2B1", "English File A2/B1 (blue) – full course"),
    "A2_Blue_Part2_Units7-12": ("EF-A2B1", "English File A2/B1 (blue) – full course"),
    "B1_Green_Course_6B-10B": ("EF-B1", "English File B1 (green) – exam course, units 6–10"),
    "Chef_Kitchen_English_Purple": ("CHEF", "Kitchen English for chefs – 10 lessons"),
}
SKILLS = ["Grammar", "Vocabulary", "Reading", "Listening", "Speaking", "Writing", "Games & songs"]
GRAMMAR = re.compile(r"simple|continuous|perfect|conditional|modal|going to|\bwill\b|relative|passive|reported|quantifier|comparative|superlative|infinitive|gerund|have to|must|should|might|\bif\b|used to|neither|question|pronoun|adverbs of manner|word order|verb be|there is|this, these|preposition|countable|some / any|how much|imperative|can /|sequencing|connector|sequencer|something|-ed|as…as|\bget\b|\bat, in, on|verb forms|say or tell|expressing movement|un- prefix|making (nouns|adjectives)|nouns? formation|compound nouns|time expressions|auxiliar", re.I)

def section(label):
    l = label.lower()
    if "answer" in l or "remember" in l or l.startswith("guarda"): return None
    if "game" in l or "song" in l or "quiz" in l: return "Games & songs"
    if l.startswith("reading"): return "Reading"
    if l.startswith("listening") or "watch" in l or "video" in l or "real tv" in l: return "Listening"
    if l.startswith("speaking") or "role-play" in l or "roleplay" in l: return "Speaking"
    if l.startswith("writing"): return "Writing"
    return None

def texts(slide):
    return [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]

def plan(folder, name):
    m = re.match(r"^[A-Za-z0-9]+_(\d{1,2}[A-C](?:-\d)?|\d{1,2}-\d{1,2}|L\d{1,2})_", name)
    if not m: return None
    raw = m.group(1)
    pid = "R" + raw if re.fullmatch(r"\d+-\d+", raw) else ("L" + str(int(raw[1:])) if raw.startswith("L") else raw)
    deck = Presentation(os.path.join(folder, name))
    first = texts(deck.slides[0])
    pages = next((t for t in first if t.startswith("Student's Book")), "")
    pages = pages.replace("Student's Book", "SB").strip()
    focus = first[2] if len(first) > 2 else ""
    skills = {}
    for s in list(deck.slides)[1:]:
        t = texts(s)
        k = section(t[2]) if len(t) > 2 else None
        if k and k not in skills: skills[k] = t[1].split("\n")[0].strip()[:80]
    if pid.startswith("R"):
        a, b = pid[1:].split("-")
        skills["Grammar"] = f"Review of units {a}–{b}"
        title = f"{first[1]} – units {a} & {b}" if first[1] == "Revise and Check" else first[1]
        skills.setdefault("Games & songs", title[:80])
        first = [first[0], title] + first[2:]
    else:
        parts = [x.strip() for x in focus.split("•")]
        g = [x for x in parts if GRAMMAR.search(x)]
        v = [x for x in parts if x and not GRAMMAR.search(x)]
        if g: skills["Grammar"] = ", ".join(g)[:80]
        if v: skills["Vocabulary"] = ", ".join(v)[:80]
        if not g and not v: skills["Vocabulary"] = first[1]
    unit = int(re.match(r"R?\d+-(\d+)|L?(\d+)", pid).group(1) or re.match(r"L?(\d+)", pid).group(1)) if pid.startswith("R") else int(re.match(r"L?(\d+)", pid).group(1))
    return {"id": pid, "unit": unit, "title": first[1], "focus": focus, "pages": pages,
            "slides": len(deck.slides), "file": f"{os.path.basename(folder)}/{name}",
            "skills": {k: skills[k] for k in SKILLS if k in skills}}

def order(p):
    m = re.match(r"(\d+)([A-C])?(?:-(\d))?", p["id"].lstrip("RL"))
    return (p["unit"], 1 if p["id"].startswith("R") else 0, m.group(2) or "", int(m.group(3) or 0))

root = sys.argv[1]
FOLDER = "https://drive.google.com/drive/folders/1_AyE9xHLZX-tR057npuZEDeoPWiVDOgB"  # "LLS Lesson Plans (teachers)"
# Drive file IDs of the uploaded decks (tools/plan-links.json: "Folder/File.pptx": id).
LINKS_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "plan-links.json")
ids = json.load(open(LINKS_FILE)) if os.path.exists(LINKS_FILE) else {}
out = {"folder": FOLDER, "links": {k: f"https://drive.google.com/file/d/{v}/view" for k, v in ids.items()}, "names": {}, "courses": {}}
for folder, (code, cname) in COURSES.items():
    path = os.path.join(root, folder)
    if not os.path.isdir(path): continue
    out["names"][code] = cname
    out["courses"].setdefault(code, [])
    for f in sorted(os.listdir(path)):
        if f.endswith(".pptx"):
            p = plan(path, f)
            if p: out["courses"][code].append(p)
for code in out["courses"]: out["courses"][code].sort(key=order)
print("/* Lesson plans (PowerPoint) per course — generated by tools/build-lesson-plans.py from the owner's decks.\n   id = the code teachers see and type in \"Unit / page\"; file = path in the slides folder. */")
print("window.LLS_PLANS = " + json.dumps(out, ensure_ascii=False, indent=1) + ";")
